#!/usr/bin/env python3
"""Koondab Rudolf Steigeri «Inimesekeskne juhtimine» TTS-teksti Wordiks."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path("/workspace")
TXT = ROOT / "Inimesekeskne-Juhtimine-TTS.txt"
DOCX = ROOT / "Inimesekeskne-Juhtimine-TTS.docx"

NAVY = RGBColor(0x0D, 0x2B, 0x4A)
TEAL = RGBColor(0x1A, 0x5F, 0x7A)
GRAY = RGBColor(0x44, 0x44, 0x44)

H1_TITLES = {
    "sisukord",
    "saateks",
    "šveitsi konföderatsiooni liidunõukogu liikme kaspar villigeri eessõna",
    "autori sissejuhatavad märkused",
}

CHAPTER_RE = re.compile(
    r"^(?:1[0-9]|20|[1-9])\. [A-ZÕÄÖÜ].{2,80}$"
)

H2_TITLES = {
    "probleemiasetus",
    "eesmärgi püstitamine",
    "tsiviil- ja sõjaväelistele juhtidele",
    "metoodilised juhised",
    "raamatu liigi kohta",
    "raamatuga töötamise kohta",
    "keskendumine olulisele",
    "peatükid",
    "märkused",
    "keelekasutus",
    "suur tänu",
    "realistlik enesehinnang",
    "investeerimine inimestesse",
    "etteantud ülesandele keskendunud ja inimesekeskne juhtimine",
    "nimede taga on inimesed",
    "erinevalt hinnatavad motiivid",
    "ausus ja avatus",
    "motivaatorite paljususe ammendamine",
    "väljakutse sõjaväejuhtidele",
    "tähelepanu inimeste pärast",
    "suunatud ja isiklik kiitus",
    "suunatud ja isiklik laitus",
    "kiituse ja laituse mitmetähenduslikkusest",
    "delikaatne kiitus",
    "virgutav laitus",
    "kriitikaga ümberkäimise kohta",
    "kontrollile esitatavad nõuded",
    "kaastöötajate kontrollimine ergutamise eesmärgil",
    "hoolitsus kaastöötajate eest",
    "infolüngad ja infotulv",
    "infovajadus ja informeerimiskohustus",
    "suuline info",
    "olla ise õigesti informeeritud",
    "kannatlik teise inimese kuulamine",
    "vahelesegamiseta teise inimese kuulamine",
    "heatahtlik teise inimese kuulamine",
    "julgustav teise inimese kuulamine",
    "vaikimine kui vastus",
    "vastuvõtt ja vestluse sissejuhatamine",
    "huvi ja osavõtt",
    "küsimuste keskne roll",
    "avatud küsimuste lai skaala",
    "suunavate küsimuste mõjust",
    "vastutulek vestluspartnerile",
    "minimaalne mõtlemisaeg",
    "ühene seisukohavõtt",
    "aus argumentatsioon",
    "kuulajasõbralik lauseehitus ja sõnavaravalik",
    "mõistmise sillad",
    "ülesaamine hirmust vestluse ees",
    "pole olemas mitte-kommunikatsiooni",
    "vestluse lõpetuseks",
    "ametikoha definitsioon ja võimalused",
    "tegemist on enama kui ainult saavutustega",
    "häired suhetes on sagedased",
    "allasurumine ei ole lahendus",
    "põgenemine konflikti eest",
    "konfliktiga leppimine",
    "võitlus konfliktipartneri vastu",
    "konfliktide tajumine ja analüüsimine",
    "konfliktide käsitlemine",
    "inimlikkus ja huumor",
    "väärtuste muutumisest ja orientatsioonikriisist tulenev väljakutse",
    "erinevad usaldusvaldkonnad",
    "iseenda usaldamine",
    "tehnika usaldamine",
    "alluvate usaldamine",
    "kolleegide ja sõprade usaldamine",
    "juhtide usaldamine",
    "usaldusväärsed juhid",
}


def is_chapter_title(line: str) -> bool:
    s = line.strip()
    if "peatükis" in s.casefold():
        return False
    return bool(CHAPTER_RE.match(s))


def is_h1(line: str, *, allow_chapters: bool) -> bool:
    s = line.strip()
    key = s.casefold()
    if key == "sisukord":
        return True
    if key in H1_TITLES:
        return allow_chapters or key == "saateks"
    return allow_chapters and is_chapter_title(s)


def is_h2(line: str) -> bool:
    return line.strip().casefold() in H2_TITLES


def add_heading(doc: Document, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    color = NAVY if level == 1 else TEAL
    for run in h.runs:
        run.font.color.rgb = color
    return h


def add_normal(doc: Document, text: str, *, italic: bool = False, bold: bool = False, center: bool = False, size: int = 12):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_after = Pt(8)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if not bold:
        run.font.color.rgb = GRAY if italic else None
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.name = "Calibri"
    return p


def load_source() -> str:
    if TXT.exists() and TXT.stat().st_size > 1000:
        return TXT.read_text(encoding="utf-8")
    raise SystemExit(f"Puudub lähtefail: {TXT}")


def build_docx(text: str) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("RUDOLF STEIGER")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = NAVY

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("INIMESEKESKNE JUHTIMINE")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY

    add_normal(
        doc,
        "Juhiseid tsiviil- ja sõjaväejuhtidele",
        italic=True,
        center=True,
        size=14,
    )
    add_normal(
        doc,
        "TTS-tekstikoond — kõik peatükid ühes failis. Kopeeri soovitud osa kõnerakendusse.",
        italic=True,
        center=True,
        size=11,
    )

    skip_title_block = True
    after_toc = False
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("#"):
            continue
        if skip_title_block:
            # skip duplicated title metadata already on cover
            if line in {
                "Rudolf Steiger",
                "INIMESEKESKNE JUHTIMINE",
                "Juhiseid tsiviil- ja sõjaväejuhtidele",
                "FONTES",
            }:
                continue
            skip_title_block = False

        if line.casefold() == "saateks":
            after_toc = True

        if is_h1(line, allow_chapters=after_toc):
            add_heading(doc, line, 1)
            continue
        if after_toc and is_h2(line):
            add_heading(doc, line, 2)
            continue
        if line.startswith("— "):
            add_bullet(doc, line[2:])
            continue
        add_normal(doc, line)

    doc.save(DOCX)
    print(f"kirjutatud {DOCX} ({DOCX.stat().st_size} baiti)")


def main() -> None:
    text = load_source()
    build_docx(text)


if __name__ == "__main__":
    main()
